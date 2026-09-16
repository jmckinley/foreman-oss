#!/usr/bin/env python3
"""Build the Foreman deck as PPTX + PDF from one slide source.

Reuses the approach the sentrygw project uses: a dark-themed editable PPTX (python-pptx) and an
HTML deck that headless Chrome prints to PDF, embedding the live dashboard screenshots in
docs/assets/. Run:  python3 scripts/build_deck.py  ->  docs/Foreman_Deck.{pptx,pdf}
"""
import base64
import os
import struct
import subprocess

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(ROOT, "docs", "assets")
OUT_DIR = os.path.join(ROOT, "docs")

# Foreman dashboard palette
BG = "0d1117"; PANEL = "151b23"; LINE = "2a333f"; TXT = "e6edf3"; MUT = "9aa7b5"
ACCENT = "7c9cf5"; GOOD = "3fb950"; WARN = "e3b341"; BAD = "f85149"; PURPLE = "a371f7"; TEAL = "39c5cf"

SLIDES = [
    {"kind": "title", "title": "Foreman",
     "tagline": "Supervise everything your agents build — and prove it’s healthy.",
     "subtitle": "The control plane for a portfolio of software built and maintained by coding "
                 "agents: it schedules the recurring reviews, collects a verdict from every run, "
                 "and surfaces what needs you. Schedules and reads — the prompts do the work.",
     "foot": "receipts in git · green/amber/red verdicts · declared-vs-effective drift · "
             "cloud / local / session tiers · one board",
     "badge": "M1–M8 shipped · 304 tests · ~86% coverage"},

    {"kind": "big", "title": "The shift",
     "big": "The bottleneck moved from writing code to supervising many autonomous streams of it.",
     "bullets": [
        "Coding agents made it cheap to create and maintain far more software per person.",
        "One operator now runs 10–30 agent-maintained projects at once.",
        "Nothing exists to keep them all healthy without babysitting each one."]},

    {"kind": "bullets", "title": "The problem", "bullets": [
        "Is each of 20 repos still documented? secure? prod-ready? on budget?",
        "CI only fires on commits — and only runs deterministic checks.",
        "The reviews that matter (architecture, UX, pen-test, prod-readiness) are open-ended "
        "judgment no linter can make.",
        "Doing it by hand, per repo, every week — doesn't scale."]},

    {"kind": "big", "title": "What Foreman is",
     "big": "A control plane: it schedules AI reviews across every project, reduces each to a "
            "durable green/amber/red verdict, and shows the fleet on one board — with a "
            "git-committed receipt behind every verdict.",
     "bullets": ["A control repo, a contract, and a small index — not a daemon, not an agent runtime.",
                 "Foreman never does the work; the cadence prompts do."]},

    {"kind": "image", "title": "One glance: catch up on the whole fleet",
     "image": "deck-01-top.png",
     "caption": "Every section folds to a one-line peek and opens on signal, so the board is a "
                "scannable digest. “How to read this board” decodes the colour language · fleet "
                "catch-up (“Since you were away”) and triage (“Needs you” · “Queued” one-click "
                "actions) lead · “Recently worked on” is a searchable per-project prompt history."},

    {"kind": "flow", "title": "How it works — every run leaves a receipt in git",
     "steps": ["a loop runs\n(cloud / local / session)", "writes a receipt\n(one JSON shape)",
               "committed to git\n(state branch)", "board reads it\n→ verdict + trend",
               "red → GitHub issue\ngreen closes it"],
     "note": "A run with no receipt is red by absence — silent scheduler failures become visible. "
             "The SQLite index is just a cache, fully rebuildable from the receipts."},

    {"kind": "image", "title": "The loop library",
     "image": "deck-04-library.png",
     "caption": "docs · quality · security · soc2 · arch · perf · prod-readiness · pen-test · "
                "test · UI/UX — one-click per project, or customize a built-in / add your own."},

    {"kind": "image", "title": "Every project’s loops, in one panel",
     "image": "deck-02-loops.png",
     "caption": "verdict + trend sparkline · flexible schedule (“every N days”, weekdays, weekly…) "
                "& next run · where it runs · run / remove / add."},

    {"kind": "image", "title": "Fleet health — with a recommendation, not just numbers",
     "image": "deck-03-repo.png",
     "caption": "Each repo leads with “Suggested next” — the raw git/GitHub stats synthesized "
                "into a short, prioritized list of what to act on (red → amber). Related signals "
                "fold into one action; the top one shows in the collapsed row. Vital-sign tiles "
                "and a cross-project roll-up sit below as detail."},

    {"kind": "big", "title": "The loop closes itself",
     "big": "amber ages to red → red opens a GitHub issue → a green run closes it, linking the "
            "clearing receipt.",
     "bullets": ["Every verdict is a schema-validated JSON receipt committed to git.",
                 "A tamper-evident, versioned record of what was reviewed, when, and how it turned out."]},

    {"kind": "pillars", "title": "Why it’s different — the moats", "pillars": [
        {"color": ACCENT, "label": "CONTRACT", "head": "Receipt-as-contract, index-as-cache",
         "desc": "Every verdict is a git-committed JSON receipt; the database is a throwaway cache. "
                 "Inspectable, versioned, portable — not a stateful black box."},
        {"color": GOOD, "label": "DRIFT", "head": "Declared ≠ effective",
         "desc": "Foreman probes what’s actually in effect (claude -p) and reports config drift — "
                 "pins, shadowed permissions, silently-skipped skills. No comparable tool does this."},
        {"color": WARN, "label": "METERED", "head": "Portfolio + time, metered",
         "desc": "Cloud/local/session tiers, per-run budgets, quota guard, spend by project — agent "
                 "runs as a metered resource across a fleet."},
        {"color": PURPLE, "label": "SUPERVISOR", "head": "A supervisor, not a swarm",
         "desc": "“Never does the work” keeps it composable and safe — it orchestrates when agents "
                 "run and reads what they concluded."}]},

    {"kind": "big", "title": "The wedge: continuous assurance",
     "big": "The security / soc2 / pen-test loops + the git-committed receipt history are, "
            "together, audit evidence that every project was reviewed on a schedule.",
     "bullets": ["Every verdict and its clearing, versioned in git.",
                 "Not a nice-to-have — audit evidence with a buyer and a budget."]},

    {"kind": "bullets", "title": "Where it sits", "bullets": [
        "Not CI/CD — cadence-driven, not commit-driven; open-ended AI reviews, not scripts.",
        "Not a code-quality SaaS — LLM judgment (arch, UX, pen-test), portfolio-wide.",
        "Not a coding agent — the layer above; it never does the work.",
        "Not an agent-swarm framework — a standing supervisor across projects and time.",
        "Uniquely: the config-drift probe, and the rebuildable git-native receipt trail."]},

    {"kind": "bullets", "title": "Who it’s for", "bullets": [
        "The operator — solo builder, studio, or platform team — running many agent-maintained projects.",
        "Who needs portfolio-level assurance: is every project still documented, secure, "
        "prod-ready, on budget — and can I prove it?"]},

    {"kind": "bullets", "title": "Status & honest limits", "bullets": [
        "Shipped: M1–M8 + secrets store + full operator surface. 304 tests, ~86% coverage.",
        "Live as a local dashboard + a token-gated read-only hosted board.",
        "Deeply wired to Claude Code today (the drift probe is the ecosystem-specific moat).",
        "Verdicts are LLM judgments — mitigated by an independent verifier pass.",
        "Single-operator today, not yet multi-tenant SaaS; cloud-tier ops partly host-bound."]},

    {"kind": "big", "title": "The bottom line",
     "big": "A supervisor you can trust and reconstruct — not an agent swarm.",
     "bullets": ["README.md · docs/ONE-PAGER.md · docs/COMPARISON.md"]},
]


def _png_size(path):
    """(w, h) from a PNG header — avoids a Pillow dependency."""
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        return None
    return struct.unpack(">II", head[16:24])


# --------------------------------------------------------------------------- PPTX (editable)

def build_pptx(slides, out_basename):
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE

    def rgb(h):
        return RGBColor.from_string(h)

    prs = Presentation()
    SW, SH = Inches(13.333), Inches(7.5)
    prs.slide_width = SW; prs.slide_height = SH
    blank = prs.slide_layouts[6]

    def bg(slide):
        r = slide.shapes.add_shape(1, 0, 0, SW, SH)
        r.fill.solid(); r.fill.fore_color.rgb = rgb(BG); r.line.fill.background()
        r.shadow.inherit = False

    def box(slide, l, t, w, h, fill=None, line=None):
        r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
        if fill:
            r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
        else:
            r.fill.background()
        if line:
            r.line.color.rgb = rgb(line); r.line.width = Pt(1)
        else:
            r.line.fill.background()
        r.shadow.inherit = False
        return r

    def text(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for i, (s, size, color, bold) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.space_after = Pt(6)
            run = p.add_run(); run.text = s
            run.font.size = Pt(size); run.font.color.rgb = rgb(color); run.font.bold = bold
            run.font.name = "Helvetica Neue"
        return tb

    for sl in slides:
        slide = prs.slides.add_slide(blank); bg(slide)
        box(slide, Inches(0.7), Inches(0.62), Inches(0.55), Inches(0.07), fill=ACCENT)
        k = sl["kind"]
        if k == "title":
            text(slide, Inches(0.9), Inches(1.6), Inches(11.5), Inches(1.2), [(sl["title"], 64, TXT, True)])
            if sl.get("tagline"):
                text(slide, Inches(0.9), Inches(2.95), Inches(11.5), Inches(1.1), [(sl["tagline"], 26, ACCENT, True)])
            text(slide, Inches(0.9), Inches(4.15), Inches(11.5), Inches(1.5), [(sl["subtitle"], 18, MUT, False)])
            if sl.get("badge"):
                b = box(slide, Inches(0.9), Inches(5.75), Inches(6.2), Inches(0.6), fill=PANEL, line=GOOD)
                tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                run = p.add_run(); run.text = sl["badge"]
                run.font.size = Pt(15); run.font.bold = True; run.font.color.rgb = rgb(GOOD)
            if sl.get("foot"):
                text(slide, Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.7), [(sl["foot"], 12, MUT, False)])
            continue

        text(slide, Inches(0.85), Inches(0.85), Inches(11.6), Inches(0.9), [(sl["title"], 32, TXT, True)])

        if k in ("bullets", "big"):
            top = Inches(2.0)
            if k == "big":
                p = box(slide, Inches(0.85), Inches(1.9), Inches(11.6), Inches(1.7), fill=PANEL, line=LINE)
                tf = p.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.3)
                run = tf.paragraphs[0].add_run(); run.text = "“" + sl["big"] + "”"
                run.font.size = Pt(22); run.font.italic = True; run.font.bold = True; run.font.color.rgb = rgb(TXT)
                top = Inches(3.9)
            tb = slide.shapes.add_textbox(Inches(0.95), top, Inches(11.4), Inches(4.2))
            tf = tb.text_frame; tf.word_wrap = True
            HANG = Inches(0.32)
            for i, b in enumerate(sl["bullets"]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(8)
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(int(HANG))); pPr.set("indent", str(-int(HANG)))
                run = p.add_run(); run.text = "• " + b
                run.font.size = Pt(17); run.font.color.rgb = rgb(TXT); run.font.name = "Helvetica Neue"

        elif k == "flow":
            steps = sl["steps"]; n = len(steps)
            gap = Inches(0.18); total = Inches(12.0)
            bw = Emu(int((total - gap * (n - 1)) / n)); x = Inches(0.7); y = Inches(2.7); bh = Inches(1.6)
            for s in steps:
                bx = box(slide, x, y, bw, bh, fill=PANEL, line=ACCENT)
                tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                for j, line in enumerate(s.split("\n")):
                    par = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    par.alignment = PP_ALIGN.CENTER; run = par.add_run(); run.text = line
                    run.font.size = Pt(13 if j == 0 else 10); run.font.bold = (j == 0)
                    run.font.color.rgb = rgb(TXT if j == 0 else MUT)
                x = Emu(int(x) + int(bw) + int(gap))
            text(slide, Inches(0.85), Inches(4.7), Inches(11.6), Inches(1.6), [(sl["note"], 16, MUT, False)])

        elif k == "image":
            cap_top = Inches(6.5)
            img = os.path.join(ASSETS, sl["image"])
            if os.path.exists(img):
                size = _png_size(img)
                if size:
                    iw, ih = size
                    maxw, maxh = Inches(11.8), Inches(3.9)
                    ratio = min(int(maxw) / iw, int(maxh) / ih)
                    w = Emu(int(iw * ratio)); h = Emu(int(ih * ratio))
                    left = Emu(int((SW - w) / 2)); top = Inches(1.55)
                    slide.shapes.add_picture(img, left, top, width=w, height=h)
                    cap_top = Emu(int(top) + int(h) + int(Inches(0.2)))
            text(slide, Inches(0.85), cap_top, Inches(11.6), Inches(1.0), [(sl["caption"], 15, MUT, False)])

        elif k == "pillars":
            y = Inches(1.95); ph = Inches(1.2); gap = Inches(0.18)
            for row in sl["pillars"]:
                c = row["color"]; header = row["label"] + " · " + row["head"]; desc = row["desc"]
                bx = box(slide, Inches(0.85), y, Inches(11.6), ph, fill=PANEL, line=c)
                tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.margin_left = Inches(0.28); tf.margin_right = Inches(0.28)
                r0 = tf.paragraphs[0].add_run(); r0.text = header
                r0.font.size = Pt(17); r0.font.bold = True; r0.font.color.rgb = rgb(c); r0.font.name = "Helvetica Neue"
                p1 = tf.add_paragraph(); p1.space_before = Pt(3); r1 = p1.add_run(); r1.text = desc
                r1.font.size = Pt(13); r1.font.color.rgb = rgb(MUT); r1.font.name = "Helvetica Neue"
                y = Emu(int(y) + int(ph) + int(gap))

    path = os.path.join(OUT_DIR, out_basename + ".pptx")
    prs.save(path)
    return path


# --------------------------------------------------------------------------- HTML -> Chrome PDF

def _esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _img_data_uri(name):
    with open(os.path.join(ASSETS, name), "rb") as f:
        return "data:image/png;base64," + base64.b64encode(f.read()).decode()


def build_html(slides, out_basename):
    css = """
    @page { size: 13.333in 7.5in; margin: 0; }
    * { box-sizing: border-box; margin: 0; padding: 0; }
    body { background:#%s; }
    .slide { width:13.333in; height:7.5in; background:#%s; color:#%s; position:relative;
      font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,Helvetica,Arial,sans-serif;
      padding:0.85in; page-break-after:always; overflow:hidden; }
    .fit { transform-origin:top left; }
    .rule { width:0.55in; height:7px; background:#%s; border-radius:3px; margin-bottom:20px;
      box-shadow:0 0 12px #%s88; }
    h1 { font-size:76px; font-weight:800; margin-top:1.0in; letter-spacing:1.5px; color:#%s; }
    .tagline { font-size:30px; font-weight:750; color:#%s; margin-top:16px; max-width:11in; line-height:1.3; }
    h2 { font-size:34px; font-weight:750; margin-bottom:22px; }
    .sub { font-size:20px; color:#%s; margin-top:14px; max-width:11in; line-height:1.45; }
    .foot { position:absolute; bottom:0.7in; left:0.85in; right:0.85in; color:#%s; font-size:14px; }
    .badge { display:inline-block; margin-top:26px; background:#%s; border:1px solid #%s; color:#%s;
      padding:11px 20px; border-radius:10px; font-weight:750; font-size:16px; }
    ul { list-style:none; margin-top:10px; }
    li { font-size:19px; line-height:1.6; margin-bottom:15px; padding-left:24px; position:relative; max-width:11.4in; }
    li:before { content:'▹'; color:#%s; position:absolute; left:0; font-weight:700; }
    .big { background:#%s; border:1px solid #%s; border-radius:14px; padding:28px 32px; font-size:25px;
      font-weight:650; font-style:italic; line-height:1.45; margin:12px 0 30px; }
    .flow { display:flex; gap:12px; margin-top:44px; }
    .step { flex:1; background:#%s; border:1px solid #%s; border-radius:12px; padding:16px 10px;
      text-align:center; min-height:1.7in; display:flex; flex-direction:column; justify-content:center; }
    .step b { font-size:15px; } .step span { font-size:12.5px; color:#%s; display:block; margin-top:6px; }
    .note { margin-top:34px; color:#%s; font-size:18px; max-width:11.4in; line-height:1.5; }
    .imgwrap { text-align:center; margin-top:0.35in; }
    .imgwrap img { max-width:11.9in; max-height:4.0in; border:1px solid #%s; border-radius:10px; }
    .caption { position:absolute; bottom:0.5in; left:0.85in; right:0.85in; color:#%s; font-size:15px; line-height:1.4; }
    .pillars { display:flex; flex-direction:column; gap:14px; margin-top:8px; }
    .prow { border-left:6px solid; border-radius:10px; background:#%s; padding:16px 22px; max-width:11.5in; }
    .prow .ph { font-size:19px; font-weight:800; margin-bottom:6px; }
    .prow .pd { font-size:16px; color:#%s; line-height:1.5; }
    """ % (BG, BG, TXT, ACCENT, ACCENT, TXT, ACCENT, MUT, MUT, PANEL, ACCENT, GOOD, ACCENT,
           PANEL, LINE, PANEL, ACCENT, MUT, MUT, LINE, MUT, PANEL, MUT)

    out = ["<!doctype html><html><head><meta charset='utf-8'><style>", css, "</style></head><body>"]
    for sl in slides:
        k = sl["kind"]
        fit = ["<div class='fit'><div class='rule'></div>"]
        pinned = []
        if k == "title":
            fit.append("<h1>%s</h1>" % _esc(sl["title"]))
            if sl.get("tagline"):
                fit.append("<div class='tagline'>%s</div>" % _esc(sl["tagline"]))
            fit.append("<div class='sub'>%s</div>" % _esc(sl["subtitle"]))
            if sl.get("badge"):
                fit.append("<div class='badge'>%s</div>" % _esc(sl["badge"]))
            if sl.get("foot"):
                pinned.append("<div class='foot'>%s</div>" % _esc(sl["foot"]))
        else:
            fit.append("<h2>%s</h2>" % _esc(sl["title"]))
            if k == "big":
                fit.append("<div class='big'>“%s”</div>" % _esc(sl["big"]))
            if k in ("bullets", "big"):
                fit.append("<ul>" + "".join("<li>%s</li>" % _esc(b) for b in sl["bullets"]) + "</ul>")
            elif k == "flow":
                fit.append("<div class='flow'>")
                for s in sl["steps"]:
                    parts = s.split("\n")
                    extra = ("<span>%s</span>" % _esc(parts[1])) if len(parts) > 1 else ""
                    fit.append("<div class='step'><b>%s</b>%s</div>" % (_esc(parts[0]), extra))
                fit.append("</div><div class='note'>%s</div>" % _esc(sl["note"]))
            elif k == "image":
                fit.append("<div class='imgwrap'><img src='%s'/></div>" % _img_data_uri(sl["image"]))
                pinned.append("<div class='caption'>%s</div>" % _esc(sl["caption"]))
            elif k == "pillars":
                fit.append("<div class='pillars'>")
                for row in sl["pillars"]:
                    fit.append("<div class='prow' style='border-left-color:#%s'>"
                               "<div class='ph' style='color:#%s'>%s · %s</div>"
                               "<div class='pd'>%s</div></div>"
                               % (row["color"], row["color"], _esc(row["label"]),
                                  _esc(row["head"]), _esc(row["desc"])))
                fit.append("</div>")
        fit.append("</div>")
        out.append("<div class='slide'>" + "".join(fit) + "".join(pinned) + "</div>")
    out.append("""<script>
window.addEventListener('load', function(){
  document.querySelectorAll('.slide').forEach(function(slide){
    var fit = slide.querySelector('.fit'); if (!fit) return;
    var cs = getComputedStyle(slide);
    var avail = slide.clientHeight - parseFloat(cs.paddingTop) - parseFloat(cs.paddingBottom);
    if (slide.querySelector('.foot, .caption')) avail -= 0.8 * 96;
    if (fit.scrollHeight > avail) {
      var k = Math.max(0.55, (avail / fit.scrollHeight) * 0.985);
      fit.style.transform = 'scale(' + k.toFixed(4) + ')';
    }
  });
});</script></body></html>""")
    path = os.path.join(OUT_DIR, "_" + out_basename + ".html")
    with open(path, "w") as f:
        f.write("".join(out))
    return path


def html_to_pdf(html_path, out_basename):
    pdf = os.path.join(OUT_DIR, out_basename + ".pdf")
    chrome = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    subprocess.run([chrome, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                    "--virtual-time-budget=5000", "--run-all-compositor-stages-before-draw",
                    "--print-to-pdf=" + pdf, "file://" + html_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return pdf


if __name__ == "__main__":
    base = "Foreman_Deck"
    print("PPTX:", build_pptx(SLIDES, base))
    html = build_html(SLIDES, base)
    print("PDF:", html_to_pdf(html, base))
    try:
        os.remove(html)
    except OSError:
        pass
